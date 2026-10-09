# -*- coding: utf-8 -*-
"""
把 extract_master.py 產出的門窗型號主檔 + batch_extract_floors.py 產出的各樓層數量，
合併輸出成矽利康用量估算 Excel（全樓層彙總 + 各樓層分頁）。

用法：
    python generate_excel.py --master master_dims.json --counts floor_counts.json \
        --out 門窗數量周長估算表.xlsx --tube-length 6
"""
import json
import argparse
import os
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "微軟正黑體"

CAT_BG = {
    "門": "FCE4D6",
    "防火門": "F8CBAD",
    "窗": "D6E4F0",
    "落地窗": "E2EFDA",
    "小窗/通風窗": "FFF2CC",
    "未分類": "D9D9D9",
}

HEADERS = ["類別", "編號", "寬(cm)", "高(cm)", "數量",
           "寬邊數(可調)", "高邊數(可調)", "單支周長(cm)", "總周長(cm)", "用量(支/管)", "備註"]
COL_W = [12, 10, 9, 9, 7, 11, 11, 12, 12, 12, 40]


def category_of(code, master):
    if code in master:
        return master[code]["category"]
    if code.startswith("SD"):
        return "防火門"
    if code.startswith("DW"):
        return "落地窗"
    if code.startswith("WB"):
        return "小窗/通風窗"
    if code.startswith("D"):
        return "門"
    if code.startswith("W"):
        return "窗"
    return "未分類"


def is_door_category(cat):
    return cat in ("門", "防火門")


def sort_key(code):
    m = re.match(r"([A-Za-z']+)(\d+)(.*)", code)
    if m:
        return (m.group(1), int(m.group(2)), m.group(3))
    return (code, 0, "")


def thin_border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)


def hdr(ws, row, col, val, bg="1F4E79", fg="FFFFFF", sz=10):
    c = ws.cell(row=row, column=col, value=val)
    c.font = Font(name=FONT_NAME, bold=True, color=fg, size=sz)
    c.fill = PatternFill("solid", fgColor=bg)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = thin_border()
    return c


def cell(ws, row, col, val, bg="FFFFFF", bold=False, color="000000", sz=10, align="center"):
    c = ws.cell(row=row, column=col, value=val)
    c.font = Font(name=FONT_NAME, bold=bold, size=sz, color=color)
    c.fill = PatternFill("solid", fgColor=bg)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    c.border = thin_border()
    return c


def formula_style(ws, row, col, bg, bold=False):
    c = ws.cell(row=row, column=col)
    c.font = Font(name=FONT_NAME, size=10, bold=bold)
    c.fill = PatternFill("solid", fgColor=bg)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = thin_border()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--master", required=True, help="extract_master.py 產出的 JSON")
    ap.add_argument("--counts", required=True, help="batch_extract_floors.py 產出的 JSON")
    ap.add_argument("--out", required=True, help="輸出 Excel 路徑")
    ap.add_argument("--tube-length", type=float, default=6,
                     help="每支(管)矽利康預設可施工長度(公尺)，Excel 內每張分頁仍可再手動調整")
    args = ap.parse_args()

    master = json.load(open(args.master, encoding="utf-8"))
    floors = json.load(open(args.counts, encoding="utf-8"))
    floor_list = list(floors.keys())
    tube_len = args.tube_length

    wb = Workbook()
    wb.remove(wb.active)
    summary_rows = {}

    for floor in floor_list:
        ws = wb.create_sheet(floor[:31])
        ws.sheet_view.showGridLines = False
        for i, w in enumerate(COL_W, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        ws.merge_cells("A1:K1")
        t = ws["A1"]
        t.value = f"{floor} 門窗數量與周長估算表"
        t.font = Font(name=FONT_NAME, bold=True, size=14, color="1F4E79")
        t.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        ws.merge_cells("A2:K2")
        s = ws["A2"]
        s.value = ("單位：cm　　周長公式：窗/落地窗/小窗＝寬×寬邊數+高×高邊數(預設2,2)；"
                    "門/防火門＝寬×寬邊數+高×高邊數(預設1,2，因底部落地不需矽利康)　　"
                    "「寬邊數/高邊數」欄可依現場狀況手動調整")
        s.font = Font(name=FONT_NAME, size=9, color="595959")
        s.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[2].height = 26

        ws.merge_cells("A3:D3")
        p = ws["A3"]
        p.value = "矽利康每支(管)可施工長度(公尺，可調)："
        p.font = Font(name=FONT_NAME, bold=True, size=10, color="1F4E79")
        p.alignment = Alignment(horizontal="right", vertical="center")
        cell(ws, 3, 5, tube_len, bg="FFFFCC", bold=True)
        ws.merge_cells("F3:K3")
        note3 = ws["F3"]
        note3.value = "用量(支/管) = ROUNDUP(總周長(cm)/100 ÷ 上方公尺數, 0)　　此數值可依實際使用之膠條/矽利康規格調整"
        note3.font = Font(name=FONT_NAME, size=9, italic=True, color="808080")
        note3.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[3].height = 20
        ws.row_dimensions[4].height = 6

        for i, h in enumerate(HEADERS, 1):
            hdr(ws, 5, i, h)

        counts = floors.get(floor, {})
        if isinstance(counts, dict) and "error" in counts:
            cell(ws, 6, 1, f"（開啟此樓層圖檔失敗：{counts['error']}）", bg="FFCCCC", align="left")
            ws.merge_cells("A6:K6")
            continue

        codes = sorted(counts.keys(), key=sort_key)
        r = 6
        for code in codes:
            cnt = counts[code]
            info = master.get(code)
            cat = category_of(code, master)
            bg = CAT_BG.get(cat, "FFFFFF")
            width_val = info["w"] if info else None
            height_val = info["h"] if info else None
            note = "" if info else "門窗主檔查無此編號尺寸資料，需人工至CAD確認寬高"

            w_mult = 1 if is_door_category(cat) else 2
            h_mult = 2

            cell(ws, r, 1, cat, bg=bg, bold=True)
            cell(ws, r, 2, code, bg=bg)
            cell(ws, r, 3, width_val, bg=bg)
            cell(ws, r, 4, height_val, bg=bg)
            cell(ws, r, 5, cnt, bg=bg, bold=True)
            cell(ws, r, 6, w_mult, bg="FFFFCC")
            cell(ws, r, 7, h_mult, bg="FFFFCC")
            if width_val is not None and height_val is not None:
                ws.cell(row=r, column=8, value=f"=C{r}*F{r}+D{r}*G{r}")
                ws.cell(row=r, column=9, value=f"=H{r}*E{r}")
                ws.cell(row=r, column=10, value=f"=ROUNDUP(I{r}/100/$E$3,0)")
            for col in (8, 9, 10):
                formula_style(ws, r, col, bg, bold=(col == 10))
            cell(ws, r, 11, note, bg="FFCCCC" if note else bg, align="left")

            entry = summary_rows.setdefault(code, {"category": cat, "w": width_val, "h": height_val,
                                                     "by_floor": {}, "note": note})
            entry["by_floor"][floor] = cnt
            r += 1

        if not codes:
            cell(ws, 6, 1, "（此樓層無門窗標籤資料）", bg="F2F2F2", align="left")
            ws.merge_cells("A6:K6")

        last_row = r - 1
        if last_row >= 6:
            tot_r = last_row + 1
            cell(ws, tot_r, 1, "小計", bg="D9E1F2", bold=True)
            for col in (2, 3, 4, 6, 7):
                cell(ws, tot_r, col, "", bg="D9E1F2")
            ws.cell(row=tot_r, column=5, value=f"=SUM(E6:E{last_row})")
            ws.cell(row=tot_r, column=9, value=f"=SUM(I6:I{last_row})")
            ws.cell(row=tot_r, column=10, value=f"=SUM(J6:J{last_row})")
            for col in (5, 8, 9, 10):
                formula_style(ws, tot_r, col, "D9E1F2", bold=True)
            cell(ws, tot_r, 8, "", bg="D9E1F2")
            cell(ws, tot_r, 11, "", bg="D9E1F2")
            ws.freeze_panes = "A6"

    # ── 全樓層彙總 ──────────────────────────
    ws_sum = wb.create_sheet("全樓層彙總", 0)
    ws_sum.sheet_view.showGridLines = False
    sum_headers = (["類別", "編號", "寬(cm)", "高(cm)", "總數量", "寬邊數(可調)", "高邊數(可調)",
                     "單支周長(cm)", "總周長(cm)", "用量(支/管)"] + floor_list + ["備註"])
    widths2 = [12, 10, 9, 9, 9, 11, 11, 12, 12, 12] + [6] * len(floor_list) + [40]
    for i, w in enumerate(widths2, 1):
        ws_sum.column_dimensions[get_column_letter(i)].width = w

    ws_sum.merge_cells(f"A1:{get_column_letter(len(sum_headers))}1")
    t = ws_sum["A1"]
    t.value = "全樓層 門窗數量與周長彙總表（矽利康施工估算用）"
    t.font = Font(name=FONT_NAME, bold=True, size=14, color="1F4E79")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 28

    ws_sum.merge_cells("A2:D2")
    p = ws_sum["A2"]
    p.value = "矽利康每支(管)可施工長度(公尺，可調)："
    p.font = Font(name=FONT_NAME, bold=True, size=10, color="1F4E79")
    p.alignment = Alignment(horizontal="right", vertical="center")
    cell(ws_sum, 2, 5, tube_len, bg="FFFFCC", bold=True)
    ws_sum.merge_cells(f"F2:{get_column_letter(len(sum_headers))}2")
    note2 = ws_sum["F2"]
    note2.value = "用量(支/管) = ROUNDUP(總周長(cm)/100 ÷ 上方公尺數, 0)　　此數值可依實際使用之膠條/矽利康規格調整"
    note2.font = Font(name=FONT_NAME, size=9, italic=True, color="808080")
    note2.alignment = Alignment(horizontal="left", vertical="center")
    ws_sum.row_dimensions[2].height = 20
    ws_sum.row_dimensions[3].height = 6

    for i, h in enumerate(sum_headers, 1):
        hdr(ws_sum, 4, i, h)

    codes_all = sorted(summary_rows.keys(), key=sort_key)
    r = 5
    for code in codes_all:
        e = summary_rows[code]
        cat = e["category"]
        bg = CAT_BG.get(cat, "FFFFFF")
        total_cnt = sum(e["by_floor"].values())
        w_mult = 1 if is_door_category(cat) else 2
        h_mult = 2

        cell(ws_sum, r, 1, cat, bg=bg, bold=True)
        cell(ws_sum, r, 2, code, bg=bg)
        cell(ws_sum, r, 3, e["w"], bg=bg)
        cell(ws_sum, r, 4, e["h"], bg=bg)
        cell(ws_sum, r, 5, total_cnt, bg=bg, bold=True)
        cell(ws_sum, r, 6, w_mult, bg="FFFFCC")
        cell(ws_sum, r, 7, h_mult, bg="FFFFCC")
        if e["w"] is not None and e["h"] is not None:
            ws_sum.cell(row=r, column=8, value=f"=C{r}*F{r}+D{r}*G{r}")
            ws_sum.cell(row=r, column=9, value=f"=H{r}*E{r}")
            ws_sum.cell(row=r, column=10, value=f"=ROUNDUP(I{r}/100/$E$2,0)")
        for col in (8, 9, 10):
            formula_style(ws_sum, r, col, bg, bold=(col == 10))

        for j, floor in enumerate(floor_list):
            col = 11 + j
            v = e["by_floor"].get(floor, "")
            cell(ws_sum, r, col, v, bg=bg)

        cell(ws_sum, r, 11 + len(floor_list), e["note"], bg="FFCCCC" if e["note"] else bg, align="left")
        r += 1

    last_row = r - 1
    if last_row >= 5:
        tot_r = last_row + 1
        cell(ws_sum, tot_r, 1, "合計", bg="D9E1F2", bold=True)
        for col in range(2, 5):
            cell(ws_sum, tot_r, col, "", bg="D9E1F2")
        ws_sum.cell(row=tot_r, column=5, value=f"=SUM(E5:E{last_row})")
        for col in (6, 7, 8):
            cell(ws_sum, tot_r, col, "", bg="D9E1F2")
        ws_sum.cell(row=tot_r, column=9, value=f"=SUM(I5:I{last_row})")
        ws_sum.cell(row=tot_r, column=10, value=f"=SUM(J5:J{last_row})")
        for col in (5, 9, 10):
            formula_style(ws_sum, tot_r, col, "D9E1F2", bold=True)
        for j, floor in enumerate(floor_list):
            col = 11 + j
            col_letter = get_column_letter(col)
            ws_sum.cell(row=tot_r, column=col, value=f"=SUM({col_letter}5:{col_letter}{last_row})")
            formula_style(ws_sum, tot_r, col, "D9E1F2", bold=True)
        cell(ws_sum, tot_r, 11 + len(floor_list), "", bg="D9E1F2")
        ws_sum.freeze_panes = "C5"

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    wb.save(args.out)
    print("Saved:", args.out)
    print("Distinct codes:", len(codes_all))


if __name__ == "__main__":
    main()
