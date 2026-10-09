import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
LABEL_FILL = PatternFill("solid", fgColor="D9EAF7")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
BAND_FILLS = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="DCE6F1")]
THIN_BORDER = Border(*(Side(style="thin", color="D9E2F3"),) * 4)
FONT_NAME = "Microsoft JhengHei"


def merge_and_set(sheet, cell_range, value, fill=None, font=None, alignment=None, border=None):
    sheet.merge_cells(cell_range)
    top_left = sheet[cell_range.split(":")[0]]
    top_left.value = value
    if fill:
        top_left.fill = fill
    if font:
        top_left.font = font
    if alignment:
        top_left.alignment = alignment
    if border:
        for row in sheet[cell_range]:
            for cell in row:
                cell.border = border


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("changes_json")
    parser.add_argument("output_xlsx")
    parser.add_argument("--old-name", default="舊版 DWG")
    parser.add_argument("--new-name", default="新版 DWG")
    args = parser.parse_args()

    data = json.loads(Path(args.changes_json).read_text(encoding="utf-8"))
    changes = data.get("changes", [])
    meta = data.get("meta", {})

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "修改彙整"
    sheet.sheet_view.showGridLines = False

    center = Alignment(horizontal="center", vertical="center")
    wrap_top = Alignment(vertical="top", wrap_text=True)
    wrap_center_left = Alignment(vertical="center", wrap_text=True)

    merge_and_set(
        sheet, "A1:F1", "DWG 圖面修改比較彙整",
        fill=TITLE_FILL, font=Font(bold=True, color="FFFFFF", size=16, name=FONT_NAME),
        alignment=center,
    )
    sheet.row_dimensions[1].height = 30

    info_rows = [
        ("舊版圖檔", args.old_name),
        ("新版圖檔", args.new_name),
        ("整理狀態", "已完成樓層／區域辨識與語意合併；備註欄標示仍需人工複核的幾何範圍。"),
    ]
    for offset, (label, value) in enumerate(info_rows):
        row = offset + 2
        cell = sheet.cell(row=row, column=1, value=label)
        cell.fill = LABEL_FILL
        cell.font = Font(bold=True, name=FONT_NAME)
        cell.alignment = center
        merge_and_set(
            sheet, f"B{row}:F{row}", value,
            font=Font(name=FONT_NAME), alignment=wrap_center_left,
        )

    headers = ["項次", "圖面／樓層", "位置", "類別", "更改內容簡述", "備註"]
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(row=6, column=column, value=header)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True, color="FFFFFF", name=FONT_NAME)
        cell.alignment = center
        cell.border = THIN_BORDER
    sheet.row_dimensions[6].height = 26

    band_index = -1
    previous_area = None
    for offset, change in enumerate(changes):
        row = 7 + offset
        area = change.get("drawing_area")
        if area != previous_area:
            band_index += 1
            previous_area = area
        band_fill = BAND_FILLS[band_index % len(BAND_FILLS)]
        values = [
            change.get("item"), change.get("drawing_area"), change.get("location"),
            change.get("category"), change.get("description"), change.get("notes") or "",
        ]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.alignment = wrap_top
            cell.border = THIN_BORDER
            cell.fill = band_fill
            if column in (1, 2, 4):
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        sheet.row_dimensions[row].height = 46

    last_row = 6 + len(changes)
    widths = [9, 18, 28, 18, 78, 30]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A7"
    if changes:
        sheet.auto_filter.ref = f"A6:F{last_row}"

    notes = workbook.create_sheet("比對說明")
    notes.sheet_view.showGridLines = False
    for column, header in enumerate(["項目", "內容"], 1):
        cell = notes.cell(row=1, column=column, value=header)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True, color="FFFFFF", name=FONT_NAME)
        cell.alignment = center
        cell.border = THIN_BORDER

    note_rows = [
        ("原始候選數量", meta.get("candidate_count", 0)),
        ("合併後交付項數", len(changes)),
        ("排除純移位數量", meta.get("pure_move_count", 0)),
        ("辨識圖面／樓層區域", meta.get("region_count", 0)),
        ("品質門檻", "通過" if meta.get("delivery_ready") else "未通過"),
        ("注意事項", "本表已經自動語意合併，但設計、施工及法規判定仍應由專業人員依核定圖說確認。"),
    ]
    for offset, (label, value) in enumerate(note_rows):
        row = offset + 2
        label_cell = notes.cell(row=row, column=1, value=label)
        label_cell.fill = LABEL_FILL
        label_cell.font = Font(bold=True, name=FONT_NAME)
        label_cell.alignment = wrap_top
        label_cell.border = THIN_BORDER
        value_cell = notes.cell(row=row, column=2, value=value)
        value_cell.font = Font(name=FONT_NAME, size=10)
        value_cell.alignment = wrap_top
        value_cell.border = THIN_BORDER

    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 92
    notes.freeze_panes = "A2"

    output_path = Path(args.output_xlsx)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    print(json.dumps({"output": str(output_path), "rows": len(changes)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
